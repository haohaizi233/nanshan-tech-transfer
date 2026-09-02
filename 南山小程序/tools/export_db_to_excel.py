# -*- coding: utf-8 -*-
"""
将 nanshan_capital 库的三张业务表导出为一个 Excel 文件（每表一个 sheet），用于交接。

用法：
  python export_db_to_excel.py
  python export_db_to_excel.py --user root --password 你的密码
  python export_db_to_excel.py --output 交接数据.xlsx

说明：
  - 默认读取 ../tools/platform-crawler/db_config.json 获取连接信息；
    其中 user/password 若为空，可用命令行参数补上。
  - 依赖：pip install pymysql openpyxl
"""
import argparse
import datetime
import json
import os

import pymysql
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_CONFIG = os.path.join(HERE, "platform-crawler", "db_config.json")

# (表名, sheet名) —— sheet 名最长 31 字符
TABLES = [
    ("platform_base", "概念验证与中试平台"),
    ("featured_invest_org", "精选投资机构"),
    ("ip_service_org", "银行融资机构"),
]

HEADER_FILL = PatternFill("solid", fgColor="2B73F8")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def load_config(path):
    with open(path, "r", encoding="utf-8") as f:
        cfg = json.load(f)
    return {
        "host": cfg.get("host", "127.0.0.1"),
        "port": int(cfg.get("port", 3306)),
        "user": cfg.get("user") or "root",
        "password": cfg.get("password") or "",
        "database": cfg.get("database", "nanshan_capital"),
    }


def fetch_table(conn, table):
    with conn.cursor() as cur:
        cur.execute(f"SELECT * FROM `{table}`")
        rows = cur.fetchall()
        columns = [d[0] for d in cur.description]
    return columns, rows


def auto_width(value):
    """按字符宽度估算列宽（中文按 2 计）。"""
    text = "" if value is None else str(value)
    width = sum(2 if ord(ch) > 127 else 1 for ch in text)
    return min(max(width, 8), 60)


def write_sheet(wb, sheet_name, columns, rows):
    ws = wb.create_sheet(title=sheet_name)
    ws.append(columns)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = "A2"
    for idx, col in enumerate(columns, start=1):
        max_w = auto_width(col)
        for row in rows:
            max_w = max(max_w, auto_width(row[idx - 1]))
        ws.column_dimensions[get_column_letter(idx)].width = max_w + 2
    return len(rows)


def main():
    parser = argparse.ArgumentParser(description="导出 nanshan_capital 三张表到 Excel")
    parser.add_argument("--config", default=DEFAULT_CONFIG, help="db_config.json 路径")
    parser.add_argument("--user", default=None, help="覆盖配置文件中的用户名")
    parser.add_argument("--password", default=None, help="覆盖配置文件中的密码")
    parser.add_argument("--output", default=None, help="输出 Excel 路径")
    args = parser.parse_args()

    cfg = load_config(args.config)
    if args.user is not None:
        cfg["user"] = args.user
    if args.password is not None:
        cfg["password"] = args.password

    stamp = datetime.datetime.now().strftime("%Y%m%d")
    output = args.output or os.path.join(HERE, f"nanshan_capital_交接数据_{stamp}.xlsx")

    conn = pymysql.connect(
        host=cfg["host"],
        port=cfg["port"],
        user=cfg["user"],
        password=cfg["password"],
        database=cfg["database"],
        charset="utf8mb4",
    )
    try:
        wb = Workbook()
        wb.remove(wb.active)
        for table, sheet in TABLES:
            columns, rows = fetch_table(conn, table)
            count = write_sheet(wb, sheet, columns, rows)
            print(f"[OK] {table} -> sheet「{sheet}」，共 {count} 行")
        wb.save(output)
    finally:
        conn.close()
    print(f"[完成] 已导出：{output}")


if __name__ == "__main__":
    main()
