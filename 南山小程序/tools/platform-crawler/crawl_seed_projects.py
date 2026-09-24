"""
南方创投网 - 种子项目库列表导出

页面:
  https://www.first-net.cn/page/portal/project/project.jsp
列表接口 (XHR POST):
  https://www.first-net.cn/portal/project_showLatestInvestmentList

用法:
  python crawl_seed_projects.py
  python crawl_seed_projects.py --out D:\\种子项目库.xlsx
"""

from __future__ import annotations

import argparse
import json
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

BASE_URL = "https://www.first-net.cn"
LIST_URL = f"{BASE_URL}/portal/project_showLatestInvestmentList"
REFERER = f"{BASE_URL}/page/portal/project/project.jsp"
PAGE_SIZE = 50

TRACK_NAMES = {
    "1": "软件及信息服务",
    "2": "集成电路",
    "3": "电子信息",
    "4": "高技术服务",
    "5": "先进制造",
    "6": "生物医疗及环保",
    "8": "新材料",
    "10": "生命健康",
    "11": "高端制造",
    "12": "电子信息",
    "13": "新能源",
    "14": "新材料",
    "15": "人工智能",
    "16": "其它",
}

_opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))


def fetch_page(page_index: int, page_size: int = PAGE_SIZE, retries: int = 3) -> dict:
    body = urllib.parse.urlencode(
        {
            "pagination.pageSize": str(page_size),
            "pagination.pageIndex": str(page_index),
            "latestInvestment.from": "pt",
            "filter": "1",
            "latestInvestment.projectTypeFlag": "1",
        }
    ).encode("utf-8")
    last_err: Exception | None = None
    for attempt in range(1, retries + 1):
        try:
            req = urllib.request.Request(
                LIST_URL,
                data=body,
                headers={
                    "User-Agent": (
                        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
                    ),
                    "X-Requested-With": "XMLHttpRequest",
                    "Origin": BASE_URL,
                    "Referer": REFERER,
                    "Accept": "application/json, text/javascript, */*; q=0.01",
                    "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
                },
            )
            with _opener.open(req, timeout=30) as resp:
                return json.loads(resp.read().decode("utf-8"))
        except (urllib.error.URLError, TimeoutError, json.JSONDecodeError) as e:
            last_err = e
            if attempt < retries:
                wait = attempt * 2
                print(f"page {page_index} failed ({attempt}/{retries}), retry in {wait}s: {e}")
                time.sleep(wait)
    raise last_err  # type: ignore[misc]


def track_name(code) -> str:
    return TRACK_NAMES.get(str(code or "").strip(), "其它")


def format_amount(amount) -> str:
    if amount is None or amount == "":
        return "0万"
    if isinstance(amount, float) and amount.is_integer():
        amount = int(amount)
    return f"{amount}万"


def detail_url(row: dict) -> str:
    project_id = (row.get("id") or "").strip()
    if not project_id:
        return ""
    if str(row.get("fromType") or "") == "2":
        return f"{BASE_URL}/page/portal/plan/introduction.jsp?planId={project_id}"
    return f"{BASE_URL}/portal/project_introduction?project.id={project_id}"


def to_excel_row(row: dict, index: int) -> list:
    referrer = (row.get("referrer") or "").strip()
    return [
        index,
        (row.get("projectName") or "").strip(),
        track_name(row.get("projectType")),
        referrer,
        format_amount(row.get("amount")),
        (row.get("companyName") or "").strip(),
        (row.get("releaseDate") or "").strip(),
        (row.get("fundingRound") or "").strip(),
        detail_url(row),
    ]


def crawl_all(page_size: int = PAGE_SIZE, sleep_sec: float = 0.25) -> tuple[int, list[dict]]:
    first = fetch_page(0, page_size)
    total = int(first.get("total") or 0)
    rows = list(first.get("result") or [])
    seen = {r.get("id") for r in rows if r.get("id")}
    print(f"total={total}, page 0 -> {len(rows)}")

    page_count = (total + page_size - 1) // page_size if total else 1
    for page_index in range(1, page_count):
        time.sleep(sleep_sec)
        data = fetch_page(page_index, page_size)
        batch = data.get("result") or []
        added = 0
        for row in batch:
            rid = row.get("id")
            if rid and rid in seen:
                continue
            if rid:
                seen.add(rid)
            rows.append(row)
            added += 1
        print(f"page {page_index}/{page_count} -> {len(batch)} new {added}, cumulative {len(rows)}")
        if not batch:
            break
    return total, rows


def save_xlsx(rows: list[dict], out_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    headers = ["序号", "项目名称", "项目赛道", "推荐人/单位", "拟融资额", "公司名称", "发布日期", "融资轮次", "详情链接"]
    wb = Workbook()
    ws = wb.active
    ws.title = "种子项目库"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for col, title in enumerate(headers, 1):
        cell = ws.cell(1, col, title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, row in enumerate(rows, 1):
        values = to_excel_row(row, i)
        for col, value in enumerate(values, 1):
            cell = ws.cell(i + 1, col, value)
            cell.alignment = Alignment(vertical="center", wrap_text=(col == 2))
            if col == 9 and value:
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")

    widths = [8, 42, 16, 36, 14, 36, 14, 12, 72]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{max(len(rows), 1) + 1}"
    ws.row_dimensions[1].height = 22

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"exported {out_path} ({len(rows)} rows)")


def main() -> None:
    parser = argparse.ArgumentParser()
    default_out = Path(__file__).resolve().parent / "种子项目库.xlsx"
    parser.add_argument("--out", type=Path, default=default_out)
    parser.add_argument("--page-size", type=int, default=PAGE_SIZE)
    parser.add_argument("--sleep", type=float, default=0.25)
    args = parser.parse_args()

    total, rows = crawl_all(args.page_size, args.sleep)
    if total and len(rows) != total:
        print(f"warning: api total={total}, collected={len(rows)}")
    save_xlsx(rows, args.out)


if __name__ == "__main__":
    main()
