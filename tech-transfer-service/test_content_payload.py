"""动态成果内容与文件上传的无数据库回归测试。"""
import base64
import unittest
from io import BytesIO
from pathlib import Path
import re
from tempfile import TemporaryDirectory

from PIL import Image

from core.errors import BizError, bad_request
from modules.achievement.content import parse_content_payload, supports_internal_patent_lookup
from modules.achievement.importer import _classify_page, _media_payload_from_uploads, _normalize_page, extract_document_records, upload_candidate_media
from modules.file.service import _image_dimensions, _validate_upload, _write_thumbnail


class ContentPayloadTests(unittest.TestCase):
    def test_parses_multiple_patents_and_ordered_sections(self):
        parsed = parse_content_payload({
            "sections": [
                {"code": "tech_problem", "title": "技术问题", "content": "要解决的问题"},
                {"code": "tech_method", "title": "技术方案", "content": "实现方案", "sortOrder": 3},
            ],
            "patents": [
                {"patentNo": "cn123", "isPrimary": True},
                {"patentNo": "CN456"},
            ],
        })
        self.assertEqual([x["code"] for x in parsed["sections"]], ["tech_problem", "tech_method"])
        self.assertEqual(parsed["patents"][0]["no"], "CN123")
        self.assertTrue(parsed["patents"][0]["primary"])

    def test_rejects_duplicate_primary_patents(self):
        with self.assertRaises(BizError):
            parse_content_payload({"patents": [
                {"patentNo": "CN1", "isPrimary": True}, {"patentNo": "CN2", "isPrimary": True},
            ]})

    def test_patent_internal_navigation_only_marks_numbers_supported_by_the_patent_module(self):
        self.assertTrue(supports_internal_patent_lookup("CN202310366203.1"))
        self.assertFalse(supports_internal_patent_lookup("内部项目编号-2026-A"))

    def test_rejects_invalid_dynamic_content_sort_orders_as_a_business_error(self):
        for payload in (
            {"sections": [{"code": "innovation", "title": "创新点", "content": "内容", "sortOrder": "bad"}]},
            {"patents": [{"patentNo": "CN1", "sortOrder": -1}]},
            {"links": [{"type": "other", "label": "链接", "url": "https://example.com", "sortOrder": True}]},
            {"media": [{"fileId": 1, "sortOrder": 10001}]},
        ):
            with self.assertRaises(BizError):
                parse_content_payload(payload)

    def test_rejects_mismatched_image_signature(self):
        with self.assertRaises(BizError):
            _validate_upload("cover.png", "image/png", b"not a png")

    def test_accepts_real_png_signature(self):
        output = BytesIO()
        Image.new("RGB", (2, 2), "white").save(output, format="PNG")
        suffix, mime = _validate_upload("cover.png", "image/png", output.getvalue())
        self.assertEqual((suffix, mime), (".png", "image/png"))

    def test_accepts_png_when_wechat_reports_generic_or_wrong_image_mime(self):
        output = BytesIO()
        Image.new("RGB", (2, 2), "white").save(output, format="PNG")
        png = output.getvalue()
        self.assertEqual(_validate_upload("cover.png", "application/octet-stream", png), (".png", "image/png"))
        self.assertEqual(_validate_upload("cover.png", "image/jpeg", png), (".png", "image/png"))

    def test_rejects_image_with_a_valid_signature_but_corrupt_image_body(self):
        with self.assertRaises(BizError):
            _validate_upload("cover.png", "image/png", b"\x89PNG\r\n\x1a\nsmall")

    def test_rejects_png_with_a_bad_idat_checksum_as_a_client_error(self):
        broken_png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAF/gJ+cyfLwQAAAABJRU5ErkJggg=="
        )
        with self.assertRaises(BizError):
            _validate_upload("cover.png", "image/png", broken_png)

    def test_legacy_media_file_ids_become_real_media_relations(self):
        parsed = parse_content_payload({"mediaFileIds": [12, "18"]})
        self.assertEqual(
            parsed["media"],
            [
                {"file_id": 12, "role": "cover", "caption": None, "alt": None, "sort": 0},
                {"file_id": 18, "role": "gallery", "caption": None, "alt": None, "sort": 1},
            ],
        )
        with self.assertRaises(BizError):
            parse_content_payload({"mediaFileIds": [12], "media": [{"fileId": 12}]})

    def test_achievement_tags_are_validated_as_a_real_relation_payload(self):
        parsed = parse_content_payload({"tags": ["人工智能", "人工智能", " AI "]})
        self.assertEqual(parsed["tags"], ["人工智能", "AI"])
        with self.assertRaises(BizError):
            parse_content_payload({"tags": [1]})

    def test_reads_uploaded_image_dimensions(self):
        output = BytesIO()
        Image.new("RGB", (640, 360), "white").save(output, format="PNG")
        self.assertEqual(_image_dimensions(".png", output.getvalue()), (640, 360))

    def test_thumbnail_writer_creates_a_bounded_jpeg_derivative(self):
        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source, thumbnail = root / "source.png", root / "thumbnail.jpg"
            Image.new("RGBA", (1600, 800), (20, 70, 180, 120)).save(source, format="PNG")
            _write_thumbnail(source, thumbnail)
            with Image.open(thumbnail) as generated:
                self.assertEqual(generated.format, "JPEG")
                self.assertLessEqual(generated.width, 320)
                self.assertLessEqual(generated.height, 320)

    def test_import_normalizes_labeled_sections_and_patents(self):
        record = _normalize_page({
            "sourcePageOrSlide": "slide-3",
            "text": "高精度视觉检测技术\n成果简介：面向工业缺陷识别。\n技术创新点：多模态融合提升小样本精度。\n技术问题：解决复杂光照误检。\n专利号：CN202310366203.1；CN202411797801.5",
        })
        normalized = record["normalized"]
        self.assertEqual(normalized["title"], "高精度视觉检测技术")
        self.assertEqual(normalized["techIntro"], "面向工业缺陷识别。")
        self.assertEqual(
            [(item["code"], item["content"]) for item in normalized["sections"]],
            [("innovation", "多模态融合提升小样本精度。"), ("tech_problem", "解决复杂光照误检。")],
        )
        self.assertEqual(
            [item["patentNo"] for item in normalized["patents"]],
            ["CN202310366203.1", "CN202411797801.5"],
        )

    def test_import_splits_contact_website_and_maturity_out_of_the_intro_blob(self):
        record = _normalize_page({
            "sourcePageOrSlide": "slide-8",
            "text": "\n".join([
                "高精度视觉检测技术",
                "面向工业缺陷识别。",
                "联系邮箱",
                "ttc@sustech.edu.cn",
                "个人网站",
                "https://ttc.sustech.edu.cn",
                "联系人：李老师",
                "电话：13900001111",
                "发明人：张三，李四",
                "成熟度：实验室验证",
            ]),
        })
        normalized = record["normalized"]
        self.assertEqual(normalized["title"], "高精度视觉检测技术")
        self.assertEqual(normalized["techIntro"], "面向工业缺陷识别。")
        self.assertEqual(normalized["contactEmail"], "ttc@sustech.edu.cn")
        self.assertEqual(normalized["contactName"], "李老师")
        self.assertEqual(normalized["contactPhone"], "13900001111")
        self.assertEqual(normalized["maturityLevel"], "02_lab")
        self.assertEqual([item["name"] for item in normalized["inventorRefs"]], ["张三", "李四"])
        self.assertEqual(normalized["links"][0]["url"], "https://ttc.sustech.edu.cn")

    def test_import_aggregates_project_and_contributor_pages_into_one_record(self):
        pages = [
            {"sourcePageOrSlide": "slide-1", "text": "基于边缘智能的航道智慧化安全管理方案\n面向工业缺陷识别。", "images": []},
            {"sourcePageOrSlide": "slide-2", "text": "成果完成人\n联系方式\n联系邮箱\nwenru1989@126.com\n文奴\n深圳职业技术大学副研究员", "images": []},
            {"sourcePageOrSlide": "slide-3", "text": "成果完成人\n联系方式\n联系邮箱\nzhang@163.com\n张三\n清华大学教授", "images": []},
        ]
        classified = []
        for page in pages:
            ptype = _classify_page(page)
            if ptype in ("empty", "template"):
                continue
            classified.append((ptype, page))
        self.assertEqual(classified[0][0], "project")
        self.assertEqual(classified[1][0], "contributor")
        self.assertEqual(classified[2][0], "contributor")

    def test_import_skips_template_placeholder_pages(self):
        self.assertEqual(_classify_page({"text": "单击添加标题", "images": []}), "template")
        self.assertEqual(_classify_page({"text": "双击此处添加文本", "images": []}), "template")
        self.assertEqual(_classify_page({"text": "Click to add title", "images": []}), "template")
        self.assertNotEqual(_classify_page({"text": "基于边缘智能的航道智慧化安全管理方案\n面向工业缺陷识别。", "images": []}), "template")

    def test_import_classifies_patent_and_contributor_pages(self):
        patent_page = {"text": "图像恢复中的滤波方法\n专利号：CN201310495188.7\n扫码查看专利公布公告", "images": []}
        self.assertEqual(_classify_page(patent_page), "patent")
        contributor_page = {"text": "成果完成人\n联系方式\n联系邮箱\nwenru1989@126.com\n文奴", "images": []}
        self.assertEqual(_classify_page(contributor_page), "contributor")

    def test_import_classifies_numbered_team_profile_as_a_contributor_page(self):
        contributor_page = {
            "text": "\n".join([
                "4 团队简介", "项目联系人及联系方式", "技术转移中心 罗老师，13927409073", "焦淑红 教授",
                "团队整体情况及核心人员介绍", "主持国家自然科学基金青年基金项目", "发表 SCI 论文 230 篇",
                "长期从事有机光电材料分子设计合成", "重点攻克光电产业材料问题", "推动产业发展",
                "冯宁，17658596827", "高级工程师", "负责研发", "获得多项专利",
            ]),
            "images": [],
        }
        self.assertEqual(_classify_page(contributor_page), "contributor")

    def test_import_classifies_bio_with_unlabelled_email_as_a_contributor_page(self):
        contributor_page = {
            "text": "成果完成人\n田雷，深圳大学特聘研究员。\n15877449448\nleitian@szu.edu.cn",
            "images": [],
        }
        self.assertEqual(_classify_page(contributor_page), "contributor")

    def test_admin_achievement_writes_keep_content_and_media(self):
        source = (Path(__file__).parent / "modules" / "admin" / "service.py").read_text(encoding="utf-8")
        self.assertIn("fields, inventor_refs, content = _parse_upsert_body", source)
        self.assertIn("sync_content(cur, ach_id", source)
        self.assertIn('"media": content.get("media") or []', source)

    def test_import_confirmation_writes_source_trace_with_the_draft_insert(self):
        base = Path(__file__).parent / "modules" / "achievement"
        service_source = (base / "service.py").read_text(encoding="utf-8")
        importer_source = (base / "importer.py").read_text(encoding="utf-8")

        self.assertRegex(service_source, r"source_import_item_id:\s*int \| None = None")
        insert_sql = service_source.split("INSERT INTO achievement", 1)[1].split("RETURNING id", 1)[0]
        self.assertIn("source_import_item_id", insert_sql)
        self.assertIn("source_import_item_id=item_id", importer_source)

    def test_import_confirmation_keeps_candidate_media_access_with_importing_user(self):
        base = Path(__file__).parent / "modules" / "achievement"
        importer_source = (base / "importer.py").read_text(encoding="utf-8")
        service_source = (base / "service.py").read_text(encoding="utf-8")

        self.assertIn("b.created_by AS imported_by_user_id", importer_source)
        self.assertIn("media_owner_user_id=int(item[\"imported_by_user_id\"])", importer_source)
        self.assertIn("media_owner_user_id: int | None = None", service_source)
        self.assertIn("media_owner_user_id or user_id", service_source)

    def test_pptx_import_extracts_embedded_images_as_candidates(self):
        from pptx import Presentation
        from pptx.util import Inches

        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            image_path = root / "candidate.png"
            Image.new("RGB", (32, 18), "navy").save(image_path, format="PNG")
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(.2), Inches(4), Inches(.5)).text_frame.text = "图像成果"
            slide.shapes.add_picture(str(image_path), Inches(1), Inches(1))
            pptx_path = root / "source.pptx"
            presentation.save(pptx_path)

            records = extract_document_records(pptx_path)

        self.assertEqual(len(records), 1)
        image = records[0]["images"][0]
        self.assertEqual(image["mimeType"], "image/png")
        self.assertEqual(image["data"][:8], b"\x89PNG\r\n\x1a\n")

    def test_pptx_import_skips_full_slide_background_images(self):
        from pptx import Presentation
        from pptx.util import Inches

        with TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            background = root / "background.png"
            content_image = root / "content.png"
            Image.new("RGB", (1600, 900), "lightblue").save(background, format="PNG")
            Image.new("RGB", (80, 50), "navy").save(content_image, format="PNG")
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_picture(str(background), 0, 0, presentation.slide_width, presentation.slide_height)
            slide.shapes.add_textbox(Inches(1), Inches(.2), Inches(5), Inches(.5)).text_frame.text = "工业视觉检测平台"
            slide.shapes.add_picture(str(content_image), Inches(1), Inches(1))
            pptx_path = root / "source.pptx"
            presentation.save(pptx_path)

            records = extract_document_records(pptx_path)

        self.assertEqual(len(records), 1)
        self.assertEqual(len(records[0]["images"]), 1)
        self.assertEqual(records[0]["images"][0]["data"][:8], b"\x89PNG\r\n\x1a\n")

    def test_import_prefers_achievement_title_over_team_section_heading(self):
        record = _normalize_page({
            "sourcePageOrSlide": "slide-3",
            "text": "\n".join([
                "团队介绍",
                "团队长期从事人工智能与零售数字化研究。",
                "面向智慧零售的多智能体协同运营平台",
                "Multi-Agent Collaborative Operation Platform",
            ]),
        })

        self.assertEqual(record["normalized"]["title"], "面向智慧零售的多智能体协同运营平台")

    def test_import_ignores_empty_template_title_field_before_later_labels(self):
        record = _normalize_page({
            "sourcePageOrSlide": "slide-1",
            "text": "\n".join([
                "扫描结构光超分辨显微镜",
                "Scanning Structured Illumination Super-Resolution Microscopy",
                "成果名称：",
                "成果简介（包括但不限于:项目背景、解决的核心问题、核心技术、研发及商业化情况等，字数800字以内）：",
                "成果照片（样机、产品)",
                "技术创新点",
                "光学成像",
            ]),
        })

        self.assertEqual(record["normalized"]["title"], "扫描结构光超分辨显微镜")
        self.assertNotIn("成果简介（包括", record["normalized"]["title"])

    def test_import_reads_numbered_project_name_and_skips_template_section_headings(self):
        record = _normalize_page({
            "sourcePageOrSlide": "slide-2",
            "text": "\n".join([
                "项目简介（200-300字）",
                "1 项目名称：面向交通运输、家庭储能等领域的新型电池",
                "项目代表性图片或照片（1-2张），或组图（1组）",
                "本项目聚焦新能源领域二次电池行业痛点。",
            ]),
        })

        self.assertEqual(record["normalized"]["title"], "面向交通运输、家庭储能等领域的新型电池")
        self.assertNotIn("项目简介（200-300字）", record["normalized"]["techIntro"])

    def test_pdf_import_extracts_embedded_images_as_candidates(self):
        with TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.pdf"
            Image.new("RGB", (30, 20), "maroon").save(source, format="PDF")
            records = extract_document_records(source)

        self.assertEqual(len(records), 1)
        image = records[0]["images"][0]
        self.assertEqual(image["mimeType"], "image/jpeg")
        self.assertEqual(image["data"][:3], b"\xff\xd8\xff")

    def test_imported_candidate_uploads_become_ordered_media_payload(self):
        media = _media_payload_from_uploads([
            {"id": 12, "name": "slide-1-image-1.png"},
            {"id": 18, "name": "slide-1-image-2.png"},
        ])
        self.assertEqual(media, [
            {"fileId": 12, "role": "cover", "sortOrder": 0},
            {"fileId": 18, "role": "gallery", "sortOrder": 1},
        ])

    def test_import_keeps_text_staged_when_one_candidate_image_is_over_limit(self):
        records = [{
            "normalized": {},
            "raw": {"text": "资料正文"},
            "images": [
                {"filename": "small.png", "mimeType": "image/png", "data": b"small"},
                {"filename": "large.png", "mimeType": "image/png", "data": b"large"},
            ],
        }]

        def uploader(_user_id, filename, _mime, _data):
            if filename == "large.png":
                raise bad_request("文件不能超过 20MB")
            return {"id": 88, "width": 32, "height": 18}

        uploaded_ids = upload_candidate_media(records, 1, uploader)

        self.assertEqual(uploaded_ids, [88])
        self.assertEqual(records[0]["normalized"]["media"], [{"fileId": 88, "role": "cover", "sortOrder": 0}])
        self.assertEqual(records[0]["raw"]["images"][0]["fileId"], 88)
        self.assertEqual(records[0]["raw"]["imageUploadWarnings"], [
            {"filename": "large.png", "message": "文件不能超过 20MB"}
        ])

    def test_import_long_unlabelled_heading_stays_staged_with_source_preserved(self):
        source_title = "团队长期从事忆阻存储及存算一体芯片研究，" * 18

        record = _normalize_page({
            "sourcePageOrSlide": "slide-28",
            "text": source_title + "\n团队负责人：杨玉超",
        })

        self.assertEqual(record["errors"], [])
        self.assertEqual(record["normalized"]["title"], source_title[:200])
        self.assertEqual(record["raw"]["sourceTitle"], source_title)
        self.assertEqual(record["raw"]["normalizationWarnings"], [{
            "field": "title",
            "message": "原始标题超过 200 字符，已截断为候选标题，请人工确认",
        }])

    def test_admin_preview_permission_covers_attached_and_staged_import_media(self):
        base = Path(__file__).parent
        service_source = (base / "modules" / "file" / "service.py").read_text(encoding="utf-8")
        route_source = (base / "routes" / "file.py").read_text(encoding="utf-8")

        self.assertIn("can_review_import_media", service_source)
        self.assertIn("achievement_media am", service_source)
        self.assertIn("organization_membership", service_source)
        self.assertIn("achievement_import_item", service_source)
        self.assertIn("raw_data @>", service_source)
        # R2-26: 文件鉴权不再信任 JWT roles 快照，改为实时查 user_role 表
        self.assertIn("SELECT 1 FROM user_role WHERE user_id", route_source)
        self.assertIn("role_code = 'admin'", route_source)

    def test_content_payload_parses_deduplicated_industry_ids(self):
        parsed = parse_content_payload({"industryIds": ["4", 9, 4]})
        self.assertEqual(parsed["industryIds"], [4, 9])

    def test_content_payload_rejects_invalid_industry_ids(self):
        with self.assertRaises(BizError):
            parse_content_payload({"industryIds": [0]})

    def test_uploaded_document_rejects_old_ppt_and_empty_payload(self):
        from modules.achievement.importer import stage_uploaded_document
        with self.assertRaises(BizError):
            stage_uploaded_document("成果.ppt", b"not-empty", 1)
        with self.assertRaises(BizError):
            stage_uploaded_document("成果.pptx", b"", 1)
        with self.assertRaises(BizError):
            stage_uploaded_document("成果.docx", b"not-empty", 1)

    def test_admin_import_upload_route_is_registered(self):
        route_source = (Path(__file__).parent / "routes" / "admin.py").read_text(encoding="utf-8")
        importer_source = (Path(__file__).parent / "modules" / "achievement" / "importer.py").read_text(encoding="utf-8")
        self.assertIn('@router.post("/admin/achievement-imports")', route_source)
        self.assertIn("stage_uploaded_documents", route_source)
        self.assertIn("list[UploadFile]", route_source)
        self.assertIn("def stage_uploaded_documents", importer_source)
        self.assertIn("旧版 .ppt 请先另存为 .pptx 后再导入", importer_source)


if __name__ == "__main__":
    unittest.main()
