import openpyxl

wb = openpyxl.load_workbook(
    r'C:\Users\admin\Desktop\南山小程序\南山小程序\tools\nanshan_capital_交接数据_20260817.xlsx',
    read_only=True, data_only=True
)
print('Sheet names:', wb.sheetnames)

for sn in wb.sheetnames:
    ws = wb[sn]
    print(f'\n=== Sheet: {sn} ===')
    print(f'Dimensions: {ws.dimensions}')
    row_count = 0
    headers = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(c) if c else '' for c in row]
            print(f'Headers ({len(headers)}):')
            for j, h in enumerate(headers):
                print(f'  [{j}] {h}')
        row_count += 1
        if i < 4 and i > 0:
            vals = []
            for c in row:
                s = str(c)[:40] if c else ''
                vals.append(s)
            print(f'Row {i}: {vals}')
    print(f'Total rows (including header): {row_count}')

wb.close()
